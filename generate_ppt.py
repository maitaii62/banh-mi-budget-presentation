from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


OUTPUT_FILE = "Ke_hoach_ngan_sach_quan_banh_mi_buoi_sang.pptx"


def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.title
    title_box.text = title
    title_box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    title_box.text_frame.paragraphs[0].font.size = Pt(24)
    title_box.text_frame.paragraphs[0].font.bold = True
    title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(30, 57, 91)

    if subtitle:
        textbox = slide.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(11), Inches(0.7))
        tf = textbox.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(90, 90, 90)


def add_bullets(slide, bullets, left=0.9, top=1.6, width=11.3, height=4.8, font_size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 5
    tf.margin_right = 5
    tf.margin_top = 5
    tf.margin_bottom = 5
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.name = "Arial"
        p.font.size = Pt(font_size)
        p.font.color.rgb = RGBColor(45, 45, 45)
        p.space_after = Pt(8)


def add_simple_table(slide, rows, title=None):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(1.5), Inches(12.1), Inches(4.6)).table
    table.columns[0].width = Inches(3.4)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(2.7)
    table.columns[3].width = Inches(3.8)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
                paragraph.font.name = "Arial"
                paragraph.font.size = Pt(12 if i == 0 else 11)
                paragraph.font.bold = (i == 0)
                paragraph.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(40, 40, 40)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(30, 57, 91) if i == 0 else RGBColor(245, 245, 245)
    if title:
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.1), Inches(8.0), Inches(0.4))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 57, 91)


def add_summary_box(slide, heading, value, subtext=None, left=0.8, top=1.5, width=3.5, height=1.8):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(230, 240, 250)
    shape.line.color.rgb = RGBColor(30, 57, 91)
    text = shape.text_frame
    p = text.paragraphs[0]
    p.text = heading
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 57, 91)
    p.alignment = PP_ALIGN.CENTER
    p2 = text.add_paragraph()
    p2.text = value
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(20, 89, 53)
    p2.alignment = PP_ALIGN.CENTER
    if subtext:
        p3 = text.add_paragraph()
        p3.text = subtext
        p3.font.size = Pt(10)
        p3.font.color.rgb = RGBColor(90, 90, 90)
        p3.alignment = PP_ALIGN.CENTER


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1 - Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(245, 248, 252))
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(23, 62, 108)
shape.line.fill.background()

logo = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(3.5), Inches(0.6))
logo_tf = logo.text_frame
logo_tf.text = "DỰ ÁN KINH DOANH"
logo_tf.paragraphs[0].font.name = "Arial"
logo_tf.paragraphs[0].font.bold = True
logo_tf.paragraphs[0].font.size = Pt(18)
logo_tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

cover_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.6), Inches(1.8))
cover_title_tf = cover_title.text_frame
p1 = cover_title_tf.paragraphs[0]
p1.text = "KẾ HOẠCH NGÂN SÁCH"
p1.font.name = "Arial"
p1.font.bold = True
p1.font.size = Pt(28)
p1.font.color.rgb = RGBColor(255, 255, 255)

p2 = cover_title_tf.add_paragraph()
p2.text = "Dự án quán bánh mì buổi sáng"
p2.font.name = "Arial"
p2.font.size = Pt(20)
p2.font.color.rgb = RGBColor(193, 214, 236)

info_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.4), Inches(5.8), Inches(1.8))
info_box.fill.solid()
info_box.fill.fore_color.rgb = RGBColor(233, 238, 246)
info_box.line.color.rgb = RGBColor(233, 238, 246)
info_text = info_box.text_frame
info_p = info_text.paragraphs[0]
info_p.text = "Sản lượng mục tiêu: 50 ổ/ngày\nDoanh thu dự kiến: 32,5 triệu/tháng\nChi phí trước dự phòng: 24 triệu/tháng"
info_p.font.name = "Arial"
info_p.font.size = Pt(16)
info_p.font.color.rgb = RGBColor(25, 58, 93)

# Slide 2 - Mô tả dự án
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "1. MÔ TẢ DỰ ÁN"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(30, 57, 91)

bullets = [
    "Tên dự án: Quán bánh mì buổi sáng – ăn nhanh, ngon và tiện lợi.",
    "Sản phẩm: bánh mì thịt, bánh mì trứng, bánh mì đặc biệt.",
    "Khách hàng mục tiêu: học sinh, sinh viên, nhân viên văn phòng, người dân quanh khu vực.",
    "Hình thức kinh doanh: bán trực tiếp tại quầy nhỏ, nhận đơn qua điện thoại và mạng xã hội.",
    "Quy mô: 1 quầy bánh mì, hoạt động chủ yếu vào buổi sáng, 50 ổ/ngày."
]
add_bullets(slide, bullets, left=0.7, top=1.5, width=11.6, height=4.5, font_size=19)

# Slide 3 - Mục tiêu ngân sách
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "2. MỤC TIÊU NGÂN SÁCH"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True

add_summary_box(slide, "Doanh thu dự kiến", "32,5 triệu", "50 ổ/ngày x 26 ngày", left=0.7, top=1.6, width=3.6, height=1.9)
add_summary_box(slide, "Chi phí trước dự phòng", "24,0 triệu", "Cố định + biến đổi", left=4.7, top=1.6, width=3.6, height=1.9)
add_summary_box(slide, "Ngân sách dự phòng", "1,2 triệu", "5% tổng chi phí", left=8.7, top=1.6, width=3.6, height=1.9)
add_summary_box(slide, "Số dư sau dự phòng", "7,3 triệu", "≈ 22,46% doanh thu", left=3.4, top=4.2, width=6.5, height=1.9)

add_bullets(slide, [
    "Mục tiêu doanh thu: 32,5 triệu/tháng.",
    "Mục tiêu kiểm soát chi phí: giữ chi phí trong ngân sách được phê duyệt.",
    "Mục tiêu lợi nhuận: đảm bảo chênh lệch thu – chi còn lại sau dự phòng."
], left=0.9, top=3.9, width=10.4, height=2.1, font_size=17)

# Slide 4 - Doanh thu dự kiến
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "3. DOANH THU DỰ KIẾN"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
rows = [
    ["Sản phẩm", "Số lượng/tháng", "Giá bán TB", "Doanh thu/tháng"],
    ["Bánh mì các loại", "1.300 ổ", "25.000đ/ổ", "32.500.000đ"],
    ["Tổng cộng", "", "", "32.500.000đ"]
]
add_simple_table(slide, rows, title="Công thức: 50 ổ/ngày x 26 ngày = 1.300 ổ/tháng")
add_bullets(slide, [
    "Doanh thu = sản lượng tiêu thụ × giá bán trung bình.",
    "Doanh thu dự kiến: 1.300 × 25.000 = 32.500.000đ/tháng."
], left=0.8, top=6.1, width=10.5, height=0.9, font_size=14)

# Slide 5 - Chi phí cố định và biến đổi
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "4. CHI PHÍ DỰ KIẾN"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True

rows1 = [
    ["Chi phí cố định", "Chi phí/tháng"],
    ["Thuê địa điểm", "3.000.000đ"],
    ["Điện, nước", "500.000đ"],
    ["Quảng cáo, in tờ rơi", "500.000đ"],
    ["Khấu hao thiết bị", "500.000đ"],
    ["Tổng chi phí cố định", "4.500.000đ"]
]
rows2 = [
    ["Chi phí biến đổi", "Chi phí/tháng"],
    ["Nguyên liệu làm bánh", "16.900.000đ"],
    ["Bao bì, túi giấy", "1.300.000đ"],
    ["Gia vị, rau và vật liệu phụ", "1.300.000đ"],
    ["Tổng chi phí biến đổi", "19.500.000đ"]
]
add_simple_table(slide, rows1)
# second table offset
shape = slide.shapes.add_table(len(rows2), len(rows2[0]), Inches(7.0), Inches(1.6), Inches(5.6), Inches(3.9)).table
for i, row in enumerate(rows2):
    for j, value in enumerate(row):
        cell = shape.cell(i, j)
        cell.text = value
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(11)
            paragraph.font.bold = (i == 0)
            paragraph.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(40, 40, 40)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 57, 91) if i == 0 else RGBColor(245, 245, 245)

# Slide 6 - Tổng hợp ngân sách
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "5. BẢNG TỔNG HỢP NGÂN SÁCH"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
rows = [
    ["Chỉ tiêu", "Giá trị"],
    ["Quy mô", "1 quầy bánh mì"],
    ["Sản lượng dự kiến", "1.300 ổ/tháng"],
    ["Giá bán trung bình", "25.000đ/ổ"],
    ["Doanh thu dự kiến", "32.500.000đ"],
    ["Tổng chi phí cố định", "4.500.000đ"],
    ["Tổng chi phí biến đổi", "19.500.000đ"],
    ["Tổng chi phí trước dự phòng", "24.000.000đ"],
    ["Dự phòng 5%", "1.200.000đ"],
    ["Tổng ngân sách bao gồm dự phòng", "25.200.000đ"]
]
add_simple_table(slide, rows)

# Slide 7 - Lợi nhuận dự kiến
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "6. LỢI NHUẬN DỰ KIẾN"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
add_summary_box(slide, "Chênh lệch trước dự phòng", "8,5 triệu", "32,5 - 24,0", left=0.8, top=1.6, width=3.8, height=1.9)
add_summary_box(slide, "Sau khi trích dự phòng", "7,3 triệu", "32,5 - 25,2", left=4.9, top=1.6, width=3.8, height=1.9)
add_summary_box(slide, "Tỷ lệ sinh lời", "22,46%", "7,3 / 32,5", left=9.0, top=1.6, width=3.2, height=1.9)
add_bullets(slide, [
    "Doanh thu dự kiến: 32,5 triệu/tháng.",
    "Chi phí trước dự phòng: 24,0 triệu/tháng.",
    "Ngân sách dự phòng: 1,2 triệu/tháng.",
    "Số tiền còn lại sau dự phòng: 7,3 triệu/tháng."
], left=0.9, top=4.2, width=11.0, height=2.2, font_size=18)

# Slide 8 - Tài trợ vốn
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "7. NGUỒN TÀI TRỢ"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
rows = [
    ["Nguồn vốn", "Số tiền"],
    ["Vốn tự có", "8.000.000đ"],
    ["Vốn góp của các thành viên", "5.000.000đ"],
    ["Vay người thân", "5.000.000đ"],
    ["Tổng vốn huy động", "18.000.000đ"]
]
add_simple_table(slide, rows)
add_bullets(slide, [
    "Vốn ban đầu đủ để chuẩn bị quầy hàng, dụng cụ và nguyên liệu ban đầu.",
    "Nhưng vốn này thấp hơn 25,2 triệu/tháng ngân sách bao gồm dự phòng, nên cần bổ sung dòng tiền."
], left=0.8, top=6.0, width=11.6, height=0.8, font_size=14)

# Slide 9 - Thời gian thực hiện
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "8. THỜI GIAN THỰC HIỆN"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
rows = [
    ["Tuần", "Công việc"],
    ["Tuần 1", "Khảo sát địa điểm, nhu cầu khách hàng, đối thủ cạnh tranh"],
    ["Tuần 2", "Tìm nhà cung cấp nguyên liệu, chuẩn bị dụng cụ"],
    ["Tuần 3", "Chuẩn bị quầy hàng, bảng hiệu, kế hoạch bán hàng"],
    ["Tuần 4", "Mua nguyên liệu, bán thử, điều chỉnh quy trình"],
    ["Tuần 5-7", "Kinh doanh chính thức, theo dõi doanh thu và chi phí"],
    ["Tuần 8", "Tổng kết, đánh giá, rút kinh nghiệm"]
]
add_simple_table(slide, rows)

# Slide 10 - Quy trình đánh giá kết quả ngân sách
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "9. QUY TRÌNH ĐÁNH GIÁ KẾT QUẢ NGÂN SÁCH"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
add_bullets(slide, [
    "Bước 1: Đánh giá chi tiêu thực tế so với kế hoạch.",
    "Bước 2: Đánh giá doanh thu và chi phí thực tế.",
    "Bước 3: Xem xét mục tiêu và chỉ tiêu ngân sách.",
    "Bước 4: Điều chỉnh ngân sách theo thực tế.",
    "Bước 5: Xác định và khắc phục thất thoát ngân sách.",
    "Bước 6: Rà soát định kỳ và rút kinh nghiệm."
], left=0.9, top=1.5, width=11.6, height=4.8, font_size=20)

# Slide 11 - Điểm hòa vốn
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "10. ĐIỂM HÒA VỐN"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
add_summary_box(slide, "Chi phí cố định", "4,5 triệu", "tháng", left=0.8, top=1.7, width=3.4, height=1.8)
add_summary_box(slide, "Giá bán TB", "25.000đ", "/ổ", left=4.7, top=1.7, width=3.4, height=1.8)
add_summary_box(slide, "Chi phí biến đổi TB", "15.000đ", "/ổ", left=8.6, top=1.7, width=3.4, height=1.8)
add_bullets(slide, [
    "Công thức: 4.500.000 ÷ (25.000 – 15.000) = 450 ổ/tháng.",
    "Sản lượng hòa vốn trung bình: 450 ÷ 26 ≈ 17,31 ổ/ngày.",
    "Như vậy, quán cần bán khoảng 18 ổ/ngày để đạt điểm hòa vốn."
], left=0.9, top=4.2, width=11.4, height=2.0, font_size=18)

# Slide 12 - Phương án cắt giảm 15%
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "11. PHƯƠNG ÁN CẮT GIẢM 15% CHI PHÍ"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
add_summary_box(slide, "Tổng chi phí cũ", "24,0 triệu", "tháng", left=0.9, top=1.7, width=3.5, height=1.8)
add_summary_box(slide, "Mức tiết kiệm", "3,6 triệu", "15%", left=5.0, top=1.7, width=3.5, height=1.8)
add_summary_box(slide, "Tổng chi phí mới", "20,4 triệu", "tháng", left=9.1, top=1.7, width=3.5, height=1.8)
add_bullets(slide, [
    "Mục tiêu: giảm tối đa khoản chi có thể kiểm soát mà vẫn giữ chất lượng sản phẩm.",
    "Ưu tiên cắt giảm: nguyên liệu, bao bì, gia vị, quảng cáo, điện nước.",
    "Không cắt giảm thuê địa điểm và khấu hao thiết bị vì là khoản chi cố định."
], left=0.9, top=4.3, width=11.5, height=1.9, font_size=18)

# Slide 13 - Cắt giảm theo mục chi
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "12. CẮT GIẢM THEO KHOẢN MỤC"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
rows = [
    ["Khoản mục", "Ban đầu", "Cắt giảm", "Sau điều chỉnh"],
    ["Nguyên liệu làm bánh", "16.900.000đ", "2.500.000đ", "14.400.000đ"],
    ["Bao bì, túi giấy", "1.300.000đ", "200.000đ", "1.100.000đ"],
    ["Gia vị, rau và vật liệu phụ", "1.300.000đ", "400.000đ", "900.000đ"],
    ["Quảng cáo, in tờ rơi", "500.000đ", "300.000đ", "200.000đ"],
    ["Điện, nước", "500.000đ", "200.000đ", "300.000đ"],
    ["Thuê địa điểm", "3.000.000đ", "0đ", "3.000.000đ"],
    ["Khấu hao thiết bị", "500.000đ", "0đ", "500.000đ"],
    ["Tổng cộng", "24.000.000đ", "3.600.000đ", "20.400.000đ"]
]
add_simple_table(slide, rows)

# Slide 14 - Dự phòng ngân sách
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "13. KẾ HOẠCH DỰ PHÒNG"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
add_bullets(slide, [
    "Có nên lập dự phòng không? Có, vì có thể phát sinh giá nguyên liệu tăng, thiết bị hỏng, doanh thu giảm.",
    "Các bước: xác định rủi ro → ước tính chi phí → chọn tỷ lệ dự phòng → xây dựng quy định sử dụng → theo dõi/điều chỉnh.",
    "Tổng chi phí trước dự phòng: 24.000.000đ/tháng; dự phòng 5% = 1.200.000đ/tháng.",
    "Tổng ngân sách bao gồm dự phòng: 25.200.000đ/tháng.",
    "Khoản dự phòng phải được quản lý riêng, chỉ sử dụng khi cần thiết."
], left=0.9, top=1.6, width=11.5, height=4.7, font_size=18)

# Slide 15 - Kiểm soát ngân sách
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(248, 249, 251))
slide.shapes.title.text = "14. KIỂM SOÁT NGÂN SÁCH SAU KHI ĐIỀU CHỈNH"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(24)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
add_bullets(slide, [
    "Ghi nhận doanh thu và chi phí phát sinh hàng ngày.",
    "Kiểm kê nguyên liệu cuối ngày để phát hiện hao hụt.",
    "So sánh giá từ các nhà cung cấp trước khi mua hàng.",
    "Đối chiếu chi phí thực tế với ngân sách hằng tuần.",
    "Nếu vượt kế hoạch, xác định nguyên nhân và điều chỉnh ngay."
], left=0.9, top=1.6, width=11.5, height=4.7, font_size=19)

# Slide 16 - Kết luận
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, RGBColor(23, 62, 108))
slide.shapes.title.text = "15. KẾT LUẬN"
slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(26)
slide.shapes.title.text_frame.paragraphs[0].font.bold = True
slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
add_bullets(slide, [
    "Dự án có tính khả thi nhờ quy mô nhỏ, địa điểm tập trung và nhu cầu tiêu thụ ổn định vào buổi sáng.",
    "Nếu doanh thu đạt 32,5 triệu/tháng và chi phí được kiểm soát đúng kế hoạch, dự án sinh lời hiệu quả.",
    "Việc lập kế hoạch dự phòng, đối chiếu ngân sách định kỳ và cắt giảm chi phí hợp lý là yếu tố then chốt để nâng cao hiệu quả kinh doanh.",
    "Kết luận: cần tăng cường kiểm soát tài chính, duy trì chất lượng sản phẩm và bảo đảm hoạt động ổn định."
], left=0.9, top=1.7, width=11.4, height=4.7, font_size=18)

# Save output
prs.save(OUTPUT_FILE)
print(f"PowerPoint generated: {OUTPUT_FILE}")
