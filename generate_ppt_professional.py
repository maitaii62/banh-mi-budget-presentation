from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUTPUT_FILE = "Ke_hoach_ngan_sach_quan_banh_mi_buoi_sang_professional.pptx"

NAVY = RGBColor(15, 42, 70)
BLUE = RGBColor(34, 79, 140)
SKY = RGBColor(93, 140, 200)
ORANGE = RGBColor(239, 123, 45)
GREEN = RGBColor(40, 149, 94)
LIGHT_BG = RGBColor(244, 247, 250)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(38, 43, 52)
MUTED = RGBColor(105, 117, 136)
LINE = RGBColor(210, 220, 232)


def add_background_slide(slide, top_color, bottom_color):
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_angle = 90
    fill.gradient_stops[0].color.rgb = top_color
    fill.gradient_stops[1].color.rgb = bottom_color


def add_title(slide, title_text, subtitle=None, align='left'):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(10.5), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Aptos'
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if align == 'center':
        p.alignment = PP_ALIGN.CENTER
    else:
        p.alignment = PP_ALIGN.LEFT

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(2.6), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = ORANGE
    line.line.color.rgb = ORANGE

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11), Inches(0.5))
        stf = sub.text_frame
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.name = 'Aptos'
        sp.font.size = Pt(11)
        sp.font.color.rgb = MUTED


def add_card(slide, left, top, width, height, label, value, suffix='', value_color=GREEN):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE
    card.line.width = Pt(1)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)

    p1 = tf.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.name = 'Aptos'
    p1.font.color.rgb = BLUE
    p1.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = f"{value}{suffix}"
    p2.font.size = Pt(23)
    p2.font.bold = True
    p2.font.name = 'Aptos'
    p2.font.color.rgb = value_color
    p2.alignment = PP_ALIGN.CENTER


def add_bullets(slide, items, left=0.9, top=1.8, width=11.3, height=4.5, font_size=18, fill_box=True):
    if fill_box:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left-0.2), Inches(top-0.2), Inches(width+0.3), Inches(height+0.3))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = LINE
        box.line.width = Pt(1)

    text_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = text_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)

    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.name = 'Aptos'
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)


def add_table(slide, rows, left, top, width, height, header_fill=NAVY, text_color=TEXT):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height))
    table = table_shape.table
    row_count = len(rows)
    col_count = len(rows[0])

    if col_count == 2:
        table.columns[0].width = Inches(width * 0.5)
        table.columns[1].width = Inches(width * 0.5)
    elif col_count == 3:
        for i in range(3):
            table.columns[i].width = Inches(width / 3)
    elif col_count == 4:
        table.columns[0].width = Inches(width * 0.28)
        table.columns[1].width = Inches(width * 0.22)
        table.columns[2].width = Inches(width * 0.24)
        table.columns[3].width = Inches(width * 0.26)

    for i in range(row_count):
        for j in range(col_count):
            cell = table.cell(i, j)
            cell.text = rows[i][j]
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = 'Aptos'
                paragraph.font.size = Pt(12 if i == 0 else 11)
                paragraph.font.color.rgb = WHITE if i == 0 else text_color
                paragraph.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
                paragraph.space_after = Pt(2)
            if i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_fill
            elif i % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(247, 250, 252)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(240, 245, 250)


def add_footer(slide):
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.08), Inches(11.8), Inches(0.2))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = 'Kế hoạch ngân sách dự án quán bánh mì buổi sáng'
    p.font.name = 'Aptos'
    p.font.size = Pt(8)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.RIGHT


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1 — Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, NAVY, BLUE)

badge = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(3.2), Inches(0.5))
pt = badge.text_frame.paragraphs[0]
pt.text = 'DỰ ÁN KINH DOANH'
pt.font.name = 'Aptos'
pt.font.size = Pt(16)
pt.font.bold = True
pt.font.color.rgb = WHITE

cover_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(9.5), Inches(1.8))
ctf = cover_title.text_frame
p1 = ctf.paragraphs[0]
p1.text = 'KẾ HOẠCH NGÂN SÁCH'
p1.font.name = 'Aptos'
p1.font.size = Pt(32)
p1.font.bold = True
p1.font.color.rgb = WHITE

p2 = ctf.add_paragraph()
p2.text = 'Dự án quán bánh mì buổi sáng'
p2.font.name = 'Aptos'
p2.font.size = Pt(22)
p2.font.color.rgb = RGBColor(200, 220, 240)

cards = [
    ('Sản lượng mục tiêu', '50', ' ổ/ngày', 0.8),
    ('Doanh thu dự kiến', '32,5', ' triệu/tháng', 4.6),
    ('Chi phí trước dự phòng', '24,0', ' triệu/tháng', 8.4),
]
for label, value, unit, left in cards:
    add_card(slide, left, 4.5, 3.1, 1.5, label, value, unit, GREEN)

# Slide 2 — Project overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '1. MÔ TẢ DỰ ÁN')
items = [
    'Tên dự án: Bánh mì buổi sáng – ăn nhanh, ngon và tiện lợi.',
    'Sản phẩm: bánh mì thịt, bánh mì trứng, bánh mì đặc biệt.',
    'Khách hàng mục tiêu: học sinh, sinh viên, nhân viên văn phòng, dân cư quanh khu vực.',
    'Hình thức kinh doanh: bán trực tiếp tại quầy; nhận đơn qua điện thoại và mạng xã hội.',
    'Quy mô: 01 quầy, hoạt động chủ yếu vào buổi sáng, mục tiêu 50 ổ/ngày.'
]
add_bullets(slide, items, left=0.9, top=1.8, width=11.2, height=4.5, font_size=18)
add_footer(slide)

# Slide 3 — Objectives
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '2. MỤC TIÊU NGÂN SÁCH')
add_card(slide, 0.9, 1.8, 2.9, 1.7, 'Doanh thu dự kiến', '32,5', ' triệu', GREEN)
add_card(slide, 4.1, 1.8, 2.9, 1.7, 'Chi phí trước dự phòng', '24,0', ' triệu', BLUE)
add_card(slide, 7.3, 1.8, 2.9, 1.7, 'Ngân sách dự phòng', '1,2', ' triệu', ORANGE)
add_card(slide, 10.5, 1.8, 2.0, 1.7, 'Lợi nhuận ước tính', '7,3', ' triệu', GREEN)

items = [
    'Mục tiêu doanh thu: 32,5 triệu/tháng từ 1.300 ổ bán ra.',
    'Mục tiêu chi phí: kiểm soát trong giới hạn ngân sách phê duyệt.',
    'Mục tiêu quản trị: giữ khoản dự phòng 1,2 triệu để ứng phó rủi ro.'
]
add_bullets(slide, items, left=0.9, top=4.2, width=11.5, height=1.8, font_size=17)
add_footer(slide)

# Slide 4 — Doanh thu
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '3. DOANH THU DỰ KIẾN')
rows = [
    ['Sản phẩm', 'Số lượng/tháng', 'Giá bán TB', 'Doanh thu/tháng'],
    ['Bánh mì các loại', '1.300 ổ', '25.000đ/ổ', '32.500.000đ'],
    ['Tổng cộng', '', '', '32.500.000đ']
]
add_table(slide, rows, 0.8, 1.7, 11.7, 2.6)
items = [
    'Công thức: 50 ổ/ngày × 26 ngày = 1.300 ổ/tháng.',
    'Doanh thu dự kiến = 1.300 × 25.000 = 32.500.000 đồng/tháng.'
]
add_bullets(slide, items, left=0.9, top=4.9, width=11.5, height=1.2, font_size=15)
add_footer(slide)

# Slide 5 — Chi phí
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '4. CHI PHÍ DỰ KIẾN')
rows1 = [
    ['Chi phí cố định', 'Chi phí/tháng'],
    ['Thuê địa điểm', '3.000.000đ'],
    ['Điện, nước', '500.000đ'],
    ['Quảng cáo, in tờ rơi', '500.000đ'],
    ['Khấu hao dụng cụ', '500.000đ'],
    ['Tổng', '4.500.000đ']
]
rows2 = [
    ['Chi phí biến đổi', 'Chi phí/tháng'],
    ['Nguyên liệu làm bánh', '16.900.000đ'],
    ['Bao bì, túi giấy', '1.300.000đ'],
    ['Gia vị, rau, vật liệu phụ', '1.300.000đ'],
    ['Tổng', '19.500.000đ']
]
add_table(slide, rows1, 0.8, 1.7, 5.7, 3.8, header_fill=NAVY)
add_table(slide, rows2, 6.8, 1.7, 5.7, 3.3, header_fill=BLUE)
add_footer(slide)

# Slide 6 — Tổng hợp ngân sách
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '5. TỔNG HỢP NGÂN SÁCH')
rows = [
    ['Chỉ tiêu', 'Giá trị'],
    ['Sản lượng dự kiến', '1.300 ổ/tháng'],
    ['Giá bán trung bình', '25.000đ/ổ'],
    ['Doanh thu dự kiến', '32.500.000đ'],
    ['Tổng chi phí cố định', '4.500.000đ'],
    ['Tổng chi phí biến đổi', '19.500.000đ'],
    ['Tổng chi phí trước dự phòng', '24.000.000đ'],
    ['Ngân sách dự phòng 5%', '1.200.000đ'],
    ['Tổng ngân sách gồm dự phòng', '25.200.000đ'],
    ['Số tiền còn lại sau dự phòng', '7.300.000đ']
]
add_table(slide, rows, 1.8, 1.7, 9.8, 5.0)
add_footer(slide)

# Slide 7 — Lợi nhuận
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '6. LỢI NHUẬN DỰ KIẾN')
add_card(slide, 0.9, 2.0, 3.5, 1.8, 'Chênh lệch trước dự phòng', '8,5', ' triệu', GREEN)
add_card(slide, 4.8, 2.0, 3.5, 1.8, 'Sau dự phòng', '7,3', ' triệu', BLUE)
add_card(slide, 8.7, 2.0, 3.6, 1.8, 'Tỷ lệ lợi nhuận', '22,46', '%', ORANGE)
items = [
    'Doanh thu 32,5 triệu – chi phí trước dự phòng 24,0 triệu = 8,5 triệu.',
    'Sau khi trích ngân sách dự phòng 1,2 triệu, số tiền còn lại là 7,3 triệu.',
    'Tỷ lệ lợi nhuận trên doanh thu ≈ 22,46%.'
]
add_bullets(slide, items, left=0.8, top=4.5, width=11.6, height=2.0, font_size=17)
add_footer(slide)

# Slide 8 — Capital source
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '7. NGUỒN TÀI TRỢ')
rows = [
    ['Nguồn vốn', 'Số tiền'],
    ['Vốn tự có', '8.000.000đ'],
    ['Vốn góp của thành viên', '5.000.000đ'],
    ['Vay người thân', '5.000.000đ'],
    ['Tổng vốn huy động', '18.000.000đ']
]
add_table(slide, rows, 2.0, 1.7, 9.3, 3.6)
items = [
    'Vốn ban đầu đủ để chuẩn bị quầy hàng, dụng cụ và mua nguyên liệu sơ bộ.',
    'Tuy nhiên, vốn 18 triệu thấp hơn tổng ngân sách 25,2 triệu/tháng nên cần có dòng tiền bổ sung.'
]
add_bullets(slide, items, left=0.9, top=5.5, width=11.4, height=1.2, font_size=15)
add_footer(slide)

# Slide 9 — Timeline
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '8. THỜI GIAN THỰC HIỆN')
rows = [
    ['Tuần', 'Công việc'],
    ['1', 'Khảo sát địa điểm, nhu cầu và đối thủ cạnh tranh'],
    ['2', 'Tìm nhà cung cấp nguyên liệu, chuẩn bị dụng cụ'],
    ['3', 'Chuẩn bị quầy hàng, bảng hiệu, kế hoạch bán hàng'],
    ['4', 'Mua nguyên liệu, bán thử và điều chỉnh quy trình'],
    ['5–7', 'Kinh doanh chính thức, theo dõi doanh thu và chi phí'],
    ['8', 'Tổng kết, đánh giá và rút kinh nghiệm']
]
add_table(slide, rows, 0.8, 1.8, 11.7, 4.4)
add_footer(slide)

# Slide 10 — Evaluation process
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '9. QUY TRÌNH ĐÁNH GIÁ KẾT QUẢ NGÂN SÁCH')
items = [
    'Bước 1: Đánh giá chi tiêu thực tế so với kế hoạch.',
    'Bước 2: Đánh giá doanh thu và chi phí thực tế.',
    'Bước 3: Xem xét mức độ hoàn thành mục tiêu ngân sách.',
    'Bước 4: Điều chỉnh ngân sách phù hợp với tình hình thực tế.',
    'Bước 5: Xác định và khắc phục thất thoát ngân sách.',
    'Bước 6: Rà soát định kỳ và rút kinh nghiệm.'
]
add_bullets(slide, items, left=0.9, top=1.9, width=11.4, height=4.3, font_size=18)
add_footer(slide)

# Slide 11 — Break-even point
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '10. ĐIỂM HÒA VỐN')
add_card(slide, 0.9, 2.1, 3.4, 1.8, 'Chi phí cố định', '4,5', ' triệu', BLUE)
add_card(slide, 4.7, 2.1, 3.6, 1.8, 'Chênh lệch giá bán - biến đổi', '10.000', 'đ/ổ', GREEN)
add_card(slide, 8.8, 2.1, 3.6, 1.8, 'Điểm hòa vốn', '450', ' ổ/tháng', ORANGE)
items = [
    'Công thức: 4.500.000 ÷ (25.000 - 15.000) = 450 ổ/tháng.',
    'Bình quân ngày: 450 ÷ 26 ≈ 17,31 ổ/ngày, khoảng 18 ổ/ngày.',
    'Kết luận: quán cần bán tối thiểu khoảng 18 ổ/ngày để không lỗ.'
]
add_bullets(slide, items, left=0.9, top=4.6, width=11.4, height=1.8, font_size=17)
add_footer(slide)

# Slide 12 — Cost reduction 15%
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '11. PHƯƠNG ÁN CẮT GIẢM 15% CHI PHÍ')
add_card(slide, 0.9, 2.0, 3.7, 1.8, 'Tổng chi phí cũ', '24,0', ' triệu', BLUE)
add_card(slide, 4.9, 2.0, 3.5, 1.8, 'Mức tiết kiệm', '3,6', ' triệu', ORANGE)
add_card(slide, 8.8, 2.0, 3.5, 1.8, 'Chi phí mới', '20,4', ' triệu', GREEN)
items = [
    'Mục tiêu: tiết kiệm 3,6 triệu/tháng, tương đương 15% chi phí trước dự phòng.',
    'Ưu tiên cắt giảm: nguyên liệu, bao bì, gia vị, quảng cáo, điện nước.',
    'Không cắt giảm chi phí thuê địa điểm và khấu hao thiết bị vì là khoản cố định.'
]
add_bullets(slide, items, left=0.9, top=4.6, width=11.4, height=1.8, font_size=17)
add_footer(slide)

# Slide 13 — Reduction by item
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '12. CẮT GIẢM CHI PHÍ THEO KHOẢN MỤC')
rows = [
    ['Khoản mục', 'Ban đầu', 'Cắt giảm', 'Sau điều chỉnh'],
    ['Nguyên liệu làm bánh', '16,9M', '2,5M', '14,4M'],
    ['Bao bì, túi giấy', '1,3M', '0,2M', '1,1M'],
    ['Gia vị, rau, vật liệu phụ', '1,3M', '0,4M', '0,9M'],
    ['Quảng cáo, in tờ rơi', '0,5M', '0,3M', '0,2M'],
    ['Điện, nước', '0,5M', '0,2M', '0,3M'],
    ['Thuê địa điểm', '3,0M', '0M', '3,0M'],
    ['Khấu hao thiết bị', '0,5M', '0M', '0,5M'],
    ['Tổng cộng', '24,0M', '3,6M', '20,4M']
]
add_table(slide, rows, 0.7, 1.7, 12.0, 4.8)
add_footer(slide)

# Slide 14 — Planning reserve
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '13. KẾ HOẠCH DỰ PHÒNG')
items = [
    'Có nên lập dự phòng không? Có, vì có thể phát sinh chi phí ngoài dự kiến như giá nguyên liệu tăng, thiết bị hư hỏng, doanh thu giảm.',
    'Các bước: xác định rủi ro → ước tính chi phí → xác định tỷ lệ dự phòng → quy định sử dụng → theo dõi và điều chỉnh.',
    'Tổng chi phí trước dự phòng: 24.000.000đ/tháng; dự phòng 5% = 1.200.000đ/tháng.',
    'Tổng ngân sách bao gồm dự phòng: 25.200.000đ/tháng.',
    'Nguyên tắc: dự phòng được quản lý riêng, chỉ sử dụng khi thực sự cần thiết.'
]
add_bullets(slide, items, left=0.9, top=1.8, width=11.3, height=4.4, font_size=17)
add_footer(slide)

# Slide 15 — Budget control
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, LIGHT_BG, WHITE)
add_title(slide, '14. KIỂM SOÁT NGÂN SÁCH')
items = [
    'Ghi nhận doanh thu và chi phí hàng ngày để theo dõi tình hình thực tế.',
    'Kiểm kê nguyên liệu cuối ngày để phát hiện hao hụt.',
    'So sánh giá từ nhiều nhà cung cấp trước khi mua hàng.',
    'Đối chiếu chi phí thực tế với ngân sách hằng tuần để phát hiện chênh lệch.',
    'Duy trì quỹ dự phòng 1.200.000đ/tháng và ghi nhận rõ việc sử dụng.',
    'Nếu vượt kế hoạch, xác định nguyên nhân và điều chỉnh kịp thời.'
]
add_bullets(slide, items, left=0.9, top=1.8, width=11.3, height=4.4, font_size=17)
add_footer(slide)

# Slide 16 — Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background_slide(slide, NAVY, BLUE)
add_title(slide, '15. KẾT LUẬN', subtitle=None)
# override title color in dark slide
for p in slide.shapes.title.text_frame.paragraphs:
    p.font.color.rgb = WHITE

items = [
    'Dự án có tính khả thi cao nhờ quy mô nhỏ, khung giờ bán hàng rõ ràng và nhu cầu tiêu thụ ổn định.',
    'Nếu doanh thu đạt 32,5 triệu/tháng và chi phí được kiểm soát tốt, dự án có thể sinh lời khoảng 7,3 triệu/tháng.',
    'Việc cắt giảm 15% chi phí và lập kế hoạch dự phòng giúp công ty chủ động hơn trước biến động giá nguyên liệu và rủi ro vận hành.',
    'Kết luận: cần duy trì kiểm soát tài chính chặt chẽ, bảo đảm chất lượng sản phẩm và quản trị ngân sách định kỳ.'
]
add_bullets(slide, items, left=0.9, top=1.8, width=11.4, height=4.2, font_size=18, fill_box=False)
for shape in slide.shapes:
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            p.font.color.rgb = WHITE

prs.save(OUTPUT_FILE)
print(f"Professional PPTX created: {OUTPUT_FILE}")
